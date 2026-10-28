from __future__ import annotations

import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from lxml import etree as ET
from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
from openpyxl.styles import Font  # type: ignore[import-untyped]

from xlinject import inject_cells, inject_cells_mixed

NS = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _create_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Measurements"
    sheet["A1"] = 1
    sheet["A1"].number_format = "0.00"
    sheet["A1"].font = Font(bold=True, color="FF123456")
    sheet["B1"] = "=A1*2"
    sheet["C1"] = "untouched"
    workbook.create_sheet("Other sheet")["A1"] = "preserve me"
    workbook.properties.description = "Synthetic preservation fixture"
    workbook.save(path)
    # Unknown archive parts must survive without being parsed or regenerated.
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("customXml/item1.xml", b'<custom xmlns="urn:example">opaque</custom>')


def _parts(path: Path) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        return {name: archive.read(name) for name in archive.namelist()}


@pytest.mark.parametrize("mixed", [False, True])
def test_highlevel_write_preserves_untargeted_parts_and_cells(tmp_path: Path, mixed: bool) -> None:
    source = tmp_path / "source with spaces \u00e4.xlsx"
    output = tmp_path / "output with spaces \u00e4.xlsx"
    _create_workbook(source)
    source_bytes = source.read_bytes()
    before = _parts(source)

    if mixed:
        report = inject_cells_mixed(
            source,
            output,
            sheet_name="Measurements",
            cell_values={"A1": "Gr\u00fc\u00dfe"},
            guard_cells=["B1"],
        )
    else:
        report = inject_cells(
            source,
            output,
            sheet_name="Measurements",
            cell_values={"A1": 12.5},
            guard_cells=["B1"],
        )

    assert report.written_count == 1
    assert source.read_bytes() == source_bytes
    after = _parts(output)
    assert before.keys() == after.keys()
    for name, contents in before.items():
        if name not in {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}:
            assert after[name] == contents, name

    old_sheet = ET.fromstring(before["xl/worksheets/sheet1.xml"])
    new_sheet = ET.fromstring(after["xl/worksheets/sheet1.xml"])
    old_cell = old_sheet.find(".//x:c[@r='A1']", NS)
    new_cell = new_sheet.find(".//x:c[@r='A1']", NS)
    assert old_cell is not None and new_cell is not None
    assert new_cell.get("s") == old_cell.get("s")
    # Once the targeted cell is removed, every other worksheet node must match.
    for cell in (old_cell, new_cell):
        parent = cell.getparent()
        assert parent is not None
        parent.remove(cell)
    assert ET.tostring(new_sheet, method="c14n") == ET.tostring(old_sheet, method="c14n")

    loaded = load_workbook(output)
    try:
        assert loaded["Measurements"]["A1"].value == ("Gr\u00fc\u00dfe" if mixed else 12.5)
        assert loaded["Measurements"]["B1"].value == "=A1*2"
        assert loaded["Measurements"]["A1"].font.bold
        assert loaded["Other sheet"]["A1"].value == "preserve me"
    finally:
        loaded.close()


def test_installed_cli_writes_and_rejects_formula_without_replacing_output(tmp_path: Path) -> None:
    source = tmp_path / "source with spaces.xlsx"
    output = tmp_path / "output with spaces.xlsx"
    _create_workbook(source)
    source_bytes = source.read_bytes()
    command = shutil.which("xlinject-write-cells")
    assert command is not None, "The package entry point must be installed"
    arguments = [
        command,
        "--input",
        str(source),
        "--output",
        str(output),
        "--sheet",
        "Measurements",
        "--guard-cells",
        "B1",
        "--cells-json",
    ]
    success = subprocess.run(
        arguments + ['{"A1":12.5}'], capture_output=True, text=True, check=True, timeout=30
    )
    assert "Written cells: 1" in success.stdout
    output_bytes = output.read_bytes()

    rejected = subprocess.run(
        arguments + ['{"B1":12.5}'], check=False, capture_output=True, text=True, timeout=30
    )
    assert rejected.returncode != 0
    assert "Refusing to overwrite formula cell" in rejected.stderr
    assert output.read_bytes() == output_bytes
    assert source.read_bytes() == source_bytes
